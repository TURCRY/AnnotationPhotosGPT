from __future__ import annotations

import ast
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "generate_word_report.py"


def _load_publish_functions():
    tree = ast.parse(SCRIPT_PATH.read_text(encoding="utf-8-sig"))
    names = {"verify_docx_package", "_unique_docx_tmp_path", "publish_docx_atomic"}
    body = []
    for node in tree.body:
        if isinstance(node, ast.Assign):
            targets = [target.id for target in node.targets if isinstance(target, ast.Name)]
            if "REQUIRED_DOCX_ENTRIES" in targets:
                body.append(node)
        elif isinstance(node, ast.FunctionDef) and node.name in names:
            body.append(node)
    module = ast.Module(body=body, type_ignores=[])
    ast.fix_missing_locations(module)
    ns = {"Path": Path, "os": os, "zipfile": zipfile, "Document": Document}
    exec(compile(module, str(SCRIPT_PATH), "exec"), ns)
    return ns


class GenerateWordReportAtomicPublishTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="annotation_word_publish_", dir=r"C:\CodexWorkspace"))
        self.ns = _load_publish_functions()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _doc(self, text="hello"):
        doc = Document()
        doc.add_paragraph(text)
        return doc

    def test_publish_normal_creates_final_and_removes_tmp(self):
        final_path = self.tmp / "report.docx"

        result = self.ns["publish_docx_atomic"](self._doc(), final_path)

        self.assertEqual(result, final_path)
        self.assertTrue(final_path.is_file())
        self.assertFalse(list(self.tmp.glob("report.docx.tmp_*")))
        self.ns["verify_docx_package"](final_path)

    def test_rename_failure_keeps_valid_tmp_and_does_not_complete(self):
        final_path = self.tmp / "report.docx"

        with mock.patch.dict(self.ns["os"].__dict__, {"replace": mock.Mock(side_effect=PermissionError("locked"))}):
            with self.assertRaises(PermissionError):
                self.ns["publish_docx_atomic"](self._doc(), final_path)

        self.assertFalse(final_path.exists())
        tmp_files = list(self.tmp.glob("report.docx.tmp_*"))
        self.assertEqual(len(tmp_files), 1)
        self.ns["verify_docx_package"](tmp_files[0])

    def test_existing_final_is_replaced_by_historical_policy(self):
        final_path = self.tmp / "report.docx"
        self._doc("old").save(final_path)

        self.ns["publish_docx_atomic"](self._doc("new"), final_path)

        reopened = Document(str(final_path))
        self.assertIn("new", "\n".join(p.text for p in reopened.paragraphs))
        self.assertNotIn("old", "\n".join(p.text for p in reopened.paragraphs))

    def test_invalid_tmp_is_not_renamed(self):
        class BadDoc:
            def save(self, path):
                Path(path).write_bytes(b"not a zip")

        final_path = self.tmp / "report.docx"

        with self.assertRaises(RuntimeError):
            self.ns["publish_docx_atomic"](BadDoc(), final_path)

        self.assertFalse(final_path.exists())
        tmp_files = list(self.tmp.glob("report.docx.tmp_*"))
        self.assertEqual(len(tmp_files), 1)


if __name__ == "__main__":
    unittest.main()
