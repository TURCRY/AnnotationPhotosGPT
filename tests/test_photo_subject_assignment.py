"""Tests de la logique metier d affectation photo -> sujet (Lot 1).

Ces tests sont independants de Streamlit et du NAS : ils n utilisent que des
fixtures locales construites a la volee dans un repertoire temporaire.
"""

import csv
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))

import photo_subject_core as core  # noqa: E402


def _write_sujets_xlsx(path: Path, header, rows, sheet_name="sujets"):
    """Ecrit un classeur Sujets.xlsx minimal via openpyxl."""
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(list(header))
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    workbook.close()


def _write_photos_batch(path: Path, names):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["photo_rel_native", "batch_status", "batch_id", "batch_ts"])
        for name in names:
            writer.writerow([name, "OK", "b1", "2026-09-01 10:00:00"])


class SujetsNormalizationTests(unittest.TestCase):
    """1. normalisation des colonnes, 2. Numero/numero, 4. numeros <= 0."""

    def test_normalize_columns_case_insensitive(self):
        mapping = core.normalize_subject_columns(["Numero", "Titre", "Localisation"])
        self.assertEqual(
            mapping, {"Numero": "numero", "Titre": "titre", "Localisation": "localisation"}
        )

    def test_normalize_columns_lowercase_and_accents(self):
        mapping = core.normalize_subject_columns(["numero", "titre", "description"])
        self.assertEqual(
            mapping, {"numero": "numero", "titre": "titre", "description": "description"}
        )

    def test_normalize_columns_ignores_unknown(self):
        mapping = core.normalize_subject_columns(["Numero", "Inconnu"])
        self.assertEqual(mapping, {"Numero": "numero"})

    def test_read_xlsx_with_capitalized_header(self):
        tmp = Path(tempfile.mkdtemp(prefix="psa_cap_", dir=r"C:\CodexWorkspace"))
        try:
            path = tmp / "Sujets.xlsx"
            _write_sujets_xlsx(
                path,
                ["Numero", "Titre", "Localisation", "Description"],
                [(1, "Sujet un", "garage", "desc"), (2, "Sujet deux", "facade", "desc")],
            )
            subjects, errors = core.read_sujets_xlsx(path)
            self.assertEqual(errors, [])
            self.assertEqual([s.numero for s in subjects], [1, 2])
            self.assertEqual(subjects[0].titre, "Sujet un")
        finally:
            _cleanup(tmp)

    def test_read_xlsx_with_lowercase_header(self):
        tmp = Path(tempfile.mkdtemp(prefix="psa_low_", dir=r"C:\CodexWorkspace"))
        try:
            path = tmp / "Sujets.xlsx"
            _write_sujets_xlsx(
                path,
                ["numero", "titre", "localisation", "description"],
                [(1, "Alpha", "loc", "d"), (2, "Beta", "loc", "d")],
            )
            subjects, errors = core.read_sujets_xlsx(path)
            self.assertEqual(errors, [])
            self.assertEqual([s.numero for s in subjects], [1, 2])
        finally:
            _cleanup(tmp)

    def test_read_xlsx_unstable_sheet_name(self):
        """Le nom de feuille n'est pas code en dur : Feuil1 doit fonctionner."""
        tmp = Path(tempfile.mkdtemp(prefix="psa_sheet_", dir=r"C:\CodexWorkspace"))
        try:
            path = tmp / "Sujets.xlsx"
            _write_sujets_xlsx(
                path,
                ["Numero", "Titre"],
                [(3, "Gamma")],
                sheet_name="Feuil1",
            )
            subjects, errors = core.read_sujets_xlsx(path)
            self.assertEqual(errors, [])
            self.assertEqual([s.numero for s in subjects], [3])
        finally:
            _cleanup(tmp)

    def test_non_positive_numbers_are_ignored(self):
        subjects, errors = core.load_subjects_from_rows(
            [
                {"numero": 0, "titre": "zero"},
                {"numero": -2, "titre": "negatif"},
                {"numero": 4, "titre": "ok"},
            ]
        )
        self.assertEqual([s.numero for s in subjects], [4])
        self.assertEqual(errors, [])

    def test_missing_numero_column_is_reported(self):
        tmp = Path(tempfile.mkdtemp(prefix="psa_nocol_", dir=r"C:\CodexWorkspace"))
        try:
            path = tmp / "Sujets.xlsx"
            _write_sujets_xlsx(path, ["Titre", "Localisation"], [("a", "b")])
            subjects, errors = core.read_sujets_xlsx(path)
            self.assertEqual(subjects, [])
            self.assertTrue(any("Numero" in e for e in errors))
        finally:
            _cleanup(tmp)


class DuplicateSubjectTests(unittest.TestCase):
    """3. detection des numeros de sujets dupliques."""

    def test_duplicate_numero_reports_error_without_resolution(self):
        subjects, errors = core.load_subjects_from_rows(
            [
                {"numero": 3, "titre": "Premier"},
                {"numero": 3, "titre": "Second"},
            ]
        )
        self.assertEqual(len(subjects), 1)
        self.assertEqual(subjects[0].titre, "Premier")
        self.assertEqual(len(errors), 1)
        self.assertIn("dupliqué", errors[0])
        self.assertIn("3", errors[0])

    def test_duplicate_does_not_use_title_to_resolve(self):
        """Aucun rapprochement par titre : le second doublon est ignore."""
        subjects, _ = core.load_subjects_from_rows(
            [
                {"numero": 7, "titre": "A"},
                {"numero": 7, "titre": "B"},
            ]
        )
        self.assertEqual([s.titre for s in subjects], ["A"])


class AssignmentsPersistenceTests(unittest.TestCase):
    """5. chargement, 6. sauvegarde, 7. ecriture atomique."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="psa_persist_", dir=r"C:\CodexWorkspace"))
        self.path = self.tmp / core.ASSIGNMENTS_FILENAME

    def tearDown(self):
        _cleanup(self.tmp)

    def test_load_missing_file_is_not_an_error(self):
        document, errors = core.load_assignments(self.tmp / "absent.json")
        self.assertEqual(errors, [])
        self.assertEqual(document.assignments, [])

    def test_save_then_load_roundtrip(self):
        document = core.AssignmentsDocument()
        subject = core.Subject(numero=3, titre="Réserve R3", localisation="garage")
        core.assign_photos(document, ["P1.JPG"], subject)
        core.save_assignments(self.path, document, id_affaire="2025-J47", id_captation="cap")

        reloaded, errors = core.load_assignments(self.path)
        self.assertEqual(errors, [])
        self.assertEqual(reloaded.id_affaire, "2025-J47")
        self.assertEqual(reloaded.id_captation, "cap")
        self.assertEqual(len(reloaded.assignments), 1)
        self.assertEqual(reloaded.assignments[0].subject_numero, 3)
        self.assertEqual(reloaded.assignments[0].subject_title, "Réserve R3")

    def test_save_preserves_accents_and_indentation(self):
        document = core.AssignmentsDocument()
        subject = core.Subject(numero=1, titre="Infiltrations d’eau", localisation="Terrasse / Sous-sol")
        core.assign_photos(document, ["P1.JPG"], subject)
        core.save_assignments(self.path, document, id_affaire="A", id_captation="C")

        raw = self.path.read_text(encoding="utf-8")
        self.assertIn("Infiltrations d’eau", raw)
        self.assertIn("\n  ", raw)  # indentation lisible

    def test_save_is_atomic_no_temp_file_left(self):
        document = core.AssignmentsDocument()
        core.exclude_photos(document, ["P1.JPG"])
        core.save_assignments(self.path, document, id_affaire="A", id_captation="C")
        leftovers = [p.name for p in self.tmp.iterdir() if p.name.startswith(".photo_subject")]
        self.assertEqual(leftovers, [])
        self.assertTrue(self.path.is_file())

    def test_save_uses_os_replace(self):
        """L'ecriture passe par un fichier temporaire puis os.replace."""
        document = core.AssignmentsDocument()
        core.exclude_photos(document, ["P1.JPG"])
        calls = []
        original_replace = os.replace
        try:
            os.replace = lambda a, b: (calls.append((str(a), str(b))), original_replace(a, b))[1]
            core.save_assignments(self.path, document, id_affaire="A", id_captation="C")
        finally:
            os.replace = original_replace
        self.assertEqual(len(calls), 1)
        self.assertTrue(calls[0][0].endswith(".tmp"))
        self.assertEqual(Path(calls[0][1]).name, core.ASSIGNMENTS_FILENAME)

    def test_non_affectee_photos_are_not_persisted(self):
        """Contrat documente : une photo non_affectee est absente du JSON."""
        document = core.AssignmentsDocument()
        core.assign_photos(document, ["P1.JPG"], core.Subject(numero=1, titre="A"))
        core.reset_photos(document, ["P1.JPG"])
        core.save_assignments(self.path, document, id_affaire="A", id_captation="C")

        data = json.loads(self.path.read_text(encoding="utf-8"))
        self.assertEqual(data["assignments"], [])
        self.assertEqual(data["schema_version"], 1)

    def test_load_rejects_unknown_status(self):
        self.path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "assignments": [{"photo_rel_native": "P1.JPG", "status": "bizarre"}],
                }
            ),
            encoding="utf-8",
        )
        document, errors = core.load_assignments(self.path)
        self.assertEqual(document.assignments, [])
        self.assertTrue(any("Statut inconnu" in e for e in errors))


class AssignmentOperationTests(unittest.TestCase):
    """8. affectation, 9. remplacement, 10. exclusion, 11. retour non_affectee."""

    def setUp(self):
        self.document = core.AssignmentsDocument()
        self.s3 = core.Subject(numero=3, titre="Réserve R3", localisation="garage")
        self.s5 = core.Subject(numero=5, titre="Réserve R5", localisation="facade")

    def test_assign_photo_to_subject(self):
        core.assign_photos(self.document, ["P1.JPG"], self.s3)
        index = core.build_assignment_index(self.document)
        self.assertEqual(core.get_status(index, "P1.JPG"), core.STATUS_AFFECTEE)
        self.assertEqual(index["P1.JPG"].subject_numero, 3)
        self.assertEqual(index["P1.JPG"].subject_title, "Réserve R3")

    def test_assign_multiple_photos_at_once(self):
        core.assign_photos(self.document, ["P1.JPG", "P2.JPG", "P3.JPG"], self.s3)
        index = core.build_assignment_index(self.document)
        self.assertEqual(
            [core.get_status(index, f"P{i}.JPG") for i in (1, 2, 3)],
            [core.STATUS_AFFECTEE] * 3,
        )

    def test_replacing_assignment_is_reported(self):
        core.assign_photos(self.document, ["P1.JPG"], self.s3)
        replaced = core.assign_photos(self.document, ["P1.JPG"], self.s5)
        index = core.build_assignment_index(self.document)
        self.assertEqual(index["P1.JPG"].subject_numero, 5)
        self.assertEqual(replaced, [("P1.JPG", "Réserve R3")])

    def test_reassign_same_subject_is_not_a_replacement(self):
        core.assign_photos(self.document, ["P1.JPG"], self.s3)
        replaced = core.assign_photos(self.document, ["P1.JPG"], self.s3)
        self.assertEqual(replaced, [])

    def test_single_primary_subject_only(self):
        core.assign_photos(self.document, ["P1.JPG"], self.s3)
        core.assign_photos(self.document, ["P1.JPG"], self.s5)
        index = core.build_assignment_index(self.document)
        self.assertEqual(len(self.document.assignments), 1)
        self.assertEqual(index["P1.JPG"].subject_numero, 5)

    def test_exclude_photo(self):
        core.assign_photos(self.document, ["P1.JPG"], self.s3)
        core.exclude_photos(self.document, ["P1.JPG"])
        index = core.build_assignment_index(self.document)
        self.assertEqual(core.get_status(index, "P1.JPG"), core.STATUS_EXCLUE)
        self.assertIsNone(index["P1.JPG"].subject_numero)
        self.assertIsNone(index["P1.JPG"].subject_title)

    def test_exclude_then_reset_to_non_affectee(self):
        core.exclude_photos(self.document, ["P1.JPG"])
        core.reset_photos(self.document, ["P1.JPG"])
        index = core.build_assignment_index(self.document)
        self.assertEqual(core.get_status(index, "P1.JPG"), core.STATUS_NON_AFFECTEE)

    def test_affected_then_reset_to_non_affectee(self):
        core.assign_photos(self.document, ["P1.JPG"], self.s3)
        core.reset_photos(self.document, ["P1.JPG"])
        index = core.build_assignment_index(self.document)
        self.assertEqual(core.get_status(index, "P1.JPG"), core.STATUS_NON_AFFECTEE)
        self.assertIsNone(index["P1.JPG"].subject_numero)


class IntegrityTests(unittest.TestCase):
    """12. photo absente, 13. sujet disparu, 14. titre/localisation modifies."""

    def setUp(self):
        self.photos = [
            core.PhotoRef(index=1, photo_rel_native="P1.JPG", display_name="P1.JPG"),
            core.PhotoRef(index=2, photo_rel_native="P2.JPG", display_name="P2.JPG"),
        ]
        self.subjects = [
            core.Subject(numero=3, titre="Réserve R3", localisation="garage"),
            core.Subject(numero=5, titre="Réserve R5", localisation="facade"),
        ]

    def test_missing_photo_is_detected(self):
        document = core.AssignmentsDocument(
            assignments=[
                core.Assignment(
                    photo_rel_native="DISPARUE.JPG",
                    status=core.STATUS_AFFECTEE,
                    subject_numero=3,
                )
            ]
        )
        issues = core.check_integrity(document, self.photos, self.subjects)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].kind, "photo_missing")
        self.assertIn("absente", issues[0].message)

    def test_missing_photo_is_not_removed(self):
        document = core.AssignmentsDocument(
            assignments=[
                core.Assignment(photo_rel_native="DISPARUE.JPG", status=core.STATUS_EXCLUE)
            ]
        )
        core.check_integrity(document, self.photos, self.subjects)
        self.assertEqual(len(document.assignments), 1)

    def test_vanished_subject_is_detected(self):
        document = core.AssignmentsDocument(
            assignments=[
                core.Assignment(
                    photo_rel_native="P1.JPG", status=core.STATUS_AFFECTEE, subject_numero=99
                )
            ]
        )
        issues = core.check_integrity(document, self.photos, self.subjects)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].kind, "subject_missing")
        self.assertIn("99", issues[0].message)

    def test_changed_title_is_detected(self):
        document = core.AssignmentsDocument(
            assignments=[
                core.Assignment(
                    photo_rel_native="P1.JPG",
                    status=core.STATUS_AFFECTEE,
                    subject_numero=3,
                    subject_title="Ancien titre",
                    subject_location="garage",
                )
            ]
        )
        issues = core.check_integrity(document, self.photos, self.subjects)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].kind, "subject_metadata_changed")
        self.assertEqual(issues[0].details["ancien_titre"], "Ancien titre")
        self.assertEqual(issues[0].details["titre_actuel"], "Réserve R3")

    def test_changed_location_is_detected(self):
        document = core.AssignmentsDocument(
            assignments=[
                core.Assignment(
                    photo_rel_native="P1.JPG",
                    status=core.STATUS_AFFECTEE,
                    subject_numero=3,
                    subject_title="Réserve R3",
                    subject_location="ancienne loc",
                )
            ]
        )
        issues = core.check_integrity(document, self.photos, self.subjects)
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].details["ancienne_localisation"], "ancienne loc")
        self.assertEqual(issues[0].details["localisation_actuelle"], "garage")

    def test_no_automatic_correction(self):
        document = core.AssignmentsDocument(
            assignments=[
                core.Assignment(
                    photo_rel_native="P1.JPG",
                    status=core.STATUS_AFFECTEE,
                    subject_numero=3,
                    subject_title="Ancien titre",
                    subject_location="garage",
                )
            ]
        )
        core.check_integrity(document, self.photos, self.subjects)
        self.assertEqual(document.assignments[0].subject_title, "Ancien titre")

    def test_excluded_photos_are_not_checked_against_subjects(self):
        document = core.AssignmentsDocument(
            assignments=[core.Assignment(photo_rel_native="P1.JPG", status=core.STATUS_EXCLUE)]
        )
        self.assertEqual(core.check_integrity(document, self.photos, self.subjects), [])


class RangeSelectionTests(unittest.TestCase):
    """15. selection d une plage suivant l ordre du CSV."""

    def setUp(self):
        # Ordre volontairement non alphabetique pour prouver qu'on utilise
        # l'ordre des lignes du CSV et non un tri de noms de fichiers.
        self.names = ["ZZ_last.JPG", "AA_first.JPG", "MM_middle.JPG", "BB_other.JPG"]
        self.photos = [
            core.PhotoRef(index=i + 1, photo_rel_native=n, display_name=n)
            for i, n in enumerate(self.names)
        ]

    def test_range_uses_csv_order_not_alphabetical(self):
        selected = core.select_range(self.photos, "AA_first.JPG", "MM_middle.JPG")
        self.assertEqual(selected, ["AA_first.JPG", "MM_middle.JPG"])

    def test_range_includes_bounds(self):
        selected = core.select_range(self.photos, "ZZ_last.JPG", "MM_middle.JPG")
        self.assertEqual(selected, ["ZZ_last.JPG", "AA_first.JPG", "MM_middle.JPG"])

    def test_reversed_bounds_still_work(self):
        selected = core.select_range(self.photos, "MM_middle.JPG", "AA_first.JPG")
        self.assertEqual(selected, ["AA_first.JPG", "MM_middle.JPG"])

    def test_single_photo_range(self):
        self.assertEqual(core.select_range(self.photos, "BB_other.JPG", "BB_other.JPG"), ["BB_other.JPG"])

    def test_unknown_bound_returns_empty(self):
        self.assertEqual(core.select_range(self.photos, "INCONNU.JPG", "AA_first.JPG"), [])

    def test_empty_bounds_return_empty(self):
        self.assertEqual(core.select_range(self.photos, None, "AA_first.JPG"), [])
        self.assertEqual(core.select_range(self.photos, "AA_first.JPG", ""), [])


class PhotosBatchOrderTests(unittest.TestCase):
    """Ordre canonique preserve a la lecture du CSV."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="psa_order_", dir=r"C:\CodexWorkspace"))

    def tearDown(self):
        _cleanup(self.tmp)

    def test_csv_row_order_is_preserved(self):
        path = self.tmp / "photos_batch.csv"
        _write_photos_batch(path, ["ZZ.JPG", "AA.JPG", "MM.JPG"])
        photos, errors = core.read_photos_batch(path)
        self.assertEqual(errors, [])
        self.assertEqual([p.photo_rel_native for p in photos], ["ZZ.JPG", "AA.JPG", "MM.JPG"])
        self.assertEqual([p.index for p in photos], [1, 2, 3])

    def test_duplicate_rows_are_ignored(self):
        path = self.tmp / "photos_batch.csv"
        _write_photos_batch(path, ["A.JPG", "A.JPG", "B.JPG"])
        photos, _ = core.read_photos_batch(path)
        self.assertEqual([p.photo_rel_native for p in photos], ["A.JPG", "B.JPG"])

    def test_missing_photo_rel_native_column(self):
        path = self.tmp / "photos_batch.csv"
        path.write_text("autre_colonne\n1\n", encoding="utf-8-sig")
        photos, errors = core.read_photos_batch(path)
        self.assertEqual(photos, [])
        self.assertTrue(any("photo_rel_native" in e for e in errors))


class GlobalFinalJoinTests(unittest.TestCase):
    """Jointure par numero avec global_final.json."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="psa_gf_", dir=r"C:\CodexWorkspace"))

    def tearDown(self):
        _cleanup(self.tmp)

    def test_numeros_are_read(self):
        path = self.tmp / "global_final.json"
        path.write_text(
            json.dumps({"sujets": [{"numero": 1}, {"numero": 3}, {"numero": 0}]}),
            encoding="utf-8",
        )
        numeros, errors = core.read_global_final_numeros(path)
        self.assertEqual(errors, [])
        self.assertEqual(numeros, {1, 3})

    def test_subject_without_cr_context_is_flagged(self):
        subjects, _ = core.load_subjects_from_rows(
            [{"numero": 1, "titre": "A"}, {"numero": 6, "titre": "B"}],
            cr_numeros={1},
        )
        self.assertTrue(subjects[0].has_cr_context)
        self.assertFalse(subjects[1].has_cr_context)

    def test_missing_global_final_is_not_blocking(self):
        numeros, errors = core.read_global_final_numeros(self.tmp / "absent.json")
        self.assertEqual(numeros, set())
        self.assertTrue(errors)


class SummaryAndFilterTests(unittest.TestCase):
    """Resume global et filtres de galerie."""

    def setUp(self):
        self.photos = [
            core.PhotoRef(index=i + 1, photo_rel_native=f"P{i}.JPG", display_name=f"P{i}.JPG")
            for i in range(1, 7)
        ]
        self.subjects = [
            core.Subject(numero=1, titre="A"),
            core.Subject(numero=2, titre="B"),
            core.Subject(numero=3, titre="C"),
        ]
        self.document = core.AssignmentsDocument()
        core.assign_photos(self.document, ["P1.JPG", "P2.JPG"], self.subjects[0])
        core.assign_photos(self.document, ["P3.JPG"], self.subjects[1])
        core.exclude_photos(self.document, ["P4.JPG"])
        self.index = core.build_assignment_index(self.document)

    def test_summary_counts(self):
        summary = core.build_summary(self.photos, self.index, self.subjects)
        self.assertEqual(summary["total"], 6)
        self.assertEqual(summary["affectees"], 3)
        self.assertEqual(summary["exclues"], 1)
        self.assertEqual(summary["non_affectees"], 2)
        self.assertEqual(summary["sujets_avec_photos"], 2)
        self.assertEqual(summary["sujets_sans_photos"], 1)

    def test_summary_per_subject(self):
        summary = core.build_summary(self.photos, self.index, self.subjects)
        self.assertEqual(summary["per_subject"], {1: 2, 2: 1, 3: 0})

    def test_filter_preserves_canonical_order(self):
        filtered = core.filter_photos(self.photos, self.index, "Affectées")
        self.assertEqual([p.photo_rel_native for p in filtered], ["P1.JPG", "P2.JPG", "P3.JPG"])

    def test_filter_non_affectees(self):
        filtered = core.filter_photos(self.photos, self.index, "Non affectées")
        self.assertEqual([p.photo_rel_native for p in filtered], ["P5.JPG", "P6.JPG"])

    def test_filter_exclues(self):
        filtered = core.filter_photos(self.photos, self.index, "Exclues")
        self.assertEqual([p.photo_rel_native for p in filtered], ["P4.JPG"])

    def test_filter_by_selected_subject(self):
        filtered = core.filter_photos(
            self.photos, self.index, "Sujet sélectionné", subject_numero=2
        )
        self.assertEqual([p.photo_rel_native for p in filtered], ["P3.JPG"])

    def test_filter_toutes(self):
        self.assertEqual(len(core.filter_photos(self.photos, self.index, "Toutes")), 6)


class PrerequisiteTests(unittest.TestCase):
    """Controle des prerequis."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="psa_prereq_", dir=r"C:\CodexWorkspace"))

    def tearDown(self):
        _cleanup(self.tmp)

    def test_all_present(self):
        sujets = self.tmp / "Sujets.xlsx"
        batch = self.tmp / "photos_batch.csv"
        gf = self.tmp / "global_final.json"
        sujets.write_text("x", encoding="utf-8")
        batch.write_text("x", encoding="utf-8")
        gf.write_text("{}", encoding="utf-8")
        checks = core.check_prerequisites(
            sujets_path=sujets,
            photos_batch_path=batch,
            global_final_path=gf,
            photos_dir=self.tmp,
            photo_count=10,
        )
        self.assertTrue(all(c["etat"] == core.STATE_OK for c in checks))

    def test_absent_files(self):
        checks = core.check_prerequisites(
            sujets_path=self.tmp / "Sujets.xlsx",
            photos_batch_path=self.tmp / "photos_batch.csv",
            global_final_path=self.tmp / "global_final.json",
            photos_dir=self.tmp,
            photo_count=0,
        )
        self.assertTrue(all(c["etat"] == core.STATE_ABSENT for c in checks))

    def test_mark_invalid(self):
        sujets = self.tmp / "Sujets.xlsx"
        sujets.write_text("x", encoding="utf-8")
        checks = core.check_prerequisites(
            sujets_path=sujets,
            photos_batch_path=None,
            global_final_path=None,
            photos_dir=self.tmp,
            photo_count=0,
        )
        core.mark_invalid(checks, "Sujets.xlsx", "colonne Numero absente")
        by_name = {c["element"]: c for c in checks}
        self.assertEqual(by_name["Sujets.xlsx"]["etat"], core.STATE_INVALID)
        self.assertEqual(by_name["Sujets.xlsx"]["detail"], "colonne Numero absente")


def _cleanup(path: Path) -> None:
    import shutil

    shutil.rmtree(path, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
