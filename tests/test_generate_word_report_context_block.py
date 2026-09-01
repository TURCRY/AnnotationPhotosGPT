from __future__ import annotations

import ast
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from docx import Document


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "generate_word_report.py"


def _load_context_block_function():
    tree = ast.parse(SCRIPT_PATH.read_text(encoding="utf-8-sig"))
    body = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "add_general_information_block"]
    module = ast.Module(body=body, type_ignores=[])
    ast.fix_missing_locations(module)
    ns = {"Document": Document, "Path": Path, "datetime": datetime}
    exec(compile(module, str(SCRIPT_PATH), "exec"), ns)
    return ns["add_general_information_block"]


class GenerateWordReportContextBlockTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="annotation_word_context_", dir=r"C:\CodexWorkspace"))
        self.add_general_information_block = _load_context_block_function()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_docx_block_uses_photo_context_not_buchelay_infos_fallback(self):
        doc = Document()
        context_path = self.tmp / "contexte_general_photos.json"
        infos = {
            "model": "modele-test",
            "mission": "Mission fallback Buchelay",
            "user": "Contexte FALLBACK BUCHELAY",
        }
        contexte = {
            "mission": "MISSION PONTOISE",
            "user": "USER PONTOISE",
            "system": "SYSTEM PONTOISE",
            "etat_avancement": "ETAT PONTOISE",
            "vlm_system": "VLM SYSTEM PONTOISE",
            "vlm_user": "VLM USER PONTOISE",
        }
        resolution = {
            "source": "contexte_general_photos",
            "source_path": str(context_path),
            "attempts": [{"level": "contexte_general_photos", "path": str(context_path), "status": "ok"}],
        }

        self.add_general_information_block(
            doc,
            annotations_path=None,
            photos_path=self.tmp / "photos.csv",
            infos=infos,
            contexte=contexte,
            context_resolution=resolution,
        )
        out = self.tmp / "context_block.docx"
        doc.save(out)
        reopened = Document(str(out))
        text = "\n".join(p.text for p in reopened.paragraphs)

        self.assertIn("MISSION PONTOISE", text)
        self.assertIn("USER PONTOISE", text)
        self.assertIn("SYSTEM PONTOISE", text)
        self.assertIn(str(context_path), text)
        self.assertNotIn("BUCHELAY", text)


if __name__ == "__main__":
    unittest.main()
